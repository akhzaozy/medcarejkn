"""Independent reconciliation of the BPJS extraction and final workbook."""
import argparse
from collections import defaultdict
import json
import math
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET
import pandas as pd
from openpyxl import load_workbook

def verify(output, source=None, workbook=False):
    data=json.loads((output/'analysis.json').read_text())
    tables={t['id']:t for t in data['tables']}
    assert len(tables)==len(data['tables'])
    for table in tables.values():
        assert table['rows']==table['processedRows'], table['id']
        if table['uniqueParticipants'] is not None:
            assert table['uniqueParticipants']<=table['rows']-table['missingParticipantId']
    distribution_totals=defaultdict(int)
    for d in data['distributions']:
        distribution_totals[(d['tableId'],d['column'])]+=d['rows']
    for (table_id,column), count in distribution_totals.items():
        assert count==tables[table_id]['rows'], (table_id,column,count)
    monthly_counts=defaultdict(int)
    monthly_costs=defaultdict(float)
    for m in data['monthly']:
        monthly_counts[m['tableId']]+=m['rows']
        for c in ['billed','verified']:
            monthly_costs[(m['tableId'],c)]+=m[c] or 0
    for table_id,count in monthly_counts.items():
        assert count==tables[table_id]['rows'], (table_id,count)
        costs=tables[table_id]['costs']
        for c,source_cols in [('billed',['FKL47','PNK17']),('verified',['FKL48','PNK18'])]:
            stats=next((costs[col] for col in source_cols if col in costs),None)
            if stats:
                assert math.isclose(monthly_costs[(table_id,c)],stats['sum'],rel_tol=1e-12,abs_tol=.01),(table_id,c)
    with zipfile.ZipFile(output/'BPJS_2025_Detail_Lengkap.zip') as archive:
        assert len(archive.namelist())==len(set(archive.namelist()))
        assert all(t['csv'] in archive.namelist() for t in tables.values())
        if source:
            assert all(str(p.relative_to(source)) in archive.namelist() for p in source.rglob('*.xlsx'))
        bad=archive.testzip()
        assert bad is None,bad
        print('ZIP CRC verified:',len(archive.namelist()),'entries',flush=True)
    if source:
        regular=next(t for t in tables.values() if t['cohort']=='Reguler' and t['kind']=='kepesertaan')
        ids=pd.read_stata(source/regular['source'],columns=['PSTV01'],convert_categoricals=False)['PSTV01']
        assert ids.nunique()==regular['uniqueParticipants']
        assert len(ids)==regular['rows']
        rtl=next(t for t in tables.values() if t['cohort']=='Reguler' and t['kind']=='fkrtl')
        costs=pd.read_stata(source/rtl['source'],columns=['FKL47','FKL48'],convert_categoricals=False)
        for c in costs:
            assert float(costs[c].sum())==rtl['costs'][c]['sum']
        print('Independent regular membership and FKRTL source totals verified',flush=True)
    if workbook:
        filename=output/'BPJS_2025_Analisis.xlsx'
        values=load_workbook(filename,read_only=True,data_only=True)
        formulas=load_workbook(filename,read_only=True,data_only=False)
        assert len(values.sheetnames)==14
        regular=next(t for t in tables.values() if t['cohort']=='Reguler' and t['kind']=='kepesertaan')
        assert values['Dashboard']['B8'].value==regular['rows']
        assert values['Dashboard']['C8'].value==regular['uniqueParticipants']
        assert values['Dashboard']['J27'].value==sum(t['rows'] for t in tables.values() if t['kind']!='referensi_icd')
        assert formulas['Dashboard']['B8'].value.startswith('=IF(')
        for kind,sheet_name in [('kepesertaan','Contoh Peserta'),('fktpkapitasi','Contoh FKTP'),('fkrtl','Contoh FKRTL'),('nonkapitasi','Contoh Nonkapitasi'),('diagnosissekunder','Contoh Diagnosis')]:
            sheet=values[sheet_name]
            headers=[c.value for c in next(sheet.iter_rows(min_row=5,max_row=5))]
            first=next(t for t in data['tables'] if t['kind']==kind)
            id_column='FKL02' if kind=='diagnosissekunder' else 'PSTV01'
            cell=sheet.cell(6,headers.index(id_column)+1)
            assert cell.value==data['samples'][first['id']][0][id_column], (sheet_name,cell.value,type(cell.value))
            assert cell.data_type=='s',(sheet_name,cell.data_type)
        for sheet in values:
            for row in sheet.iter_rows():
                for cell in row:
                    assert cell.data_type!='e',(sheet.title,cell.coordinate,cell.value)
        values.close(); formulas.close()
        with zipfile.ZipFile(filename) as z:
            chart_names=[n for n in z.namelist() if '/charts/chart' in n and n.endswith('.xml')]
            assert len(chart_names)==1
            chart=z.read(chart_names[0]).decode()
            assert '007A78' in chart.upper()
            assert '$A$24:$A$30' in chart and '$B$24:$B$30' in chart
            xml=z.read('xl/worksheets/sheet1.xml').decode()
            assert 'dataValidation' in xml and 'Reguler' in xml
            ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
            pane=ET.fromstring(z.read('xl/worksheets/sheet2.xml')).find('.//m:pane',ns)
            assert pane.attrib['xSplit']=='2' and pane.attrib['ySplit']=='5'
            assert sum(n.startswith('xl/tables/table') and n.endswith('.xml') for n in z.namelist())==13
        print('Workbook values, formulas, chart, selector, tables, freeze panes verified',flush=True)
    print('Verified:',len(tables),'Stata tables;',sum(t['rows'] for t in tables.values() if t['kind']!='referensi_icd'),'main rows;',len(distribution_totals),'distributions',flush=True)

if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--source',type=Path)
    parser.add_argument('--workbook',action='store_true')
    args=parser.parse_args()
    verify(args.output,args.source,args.workbook)
